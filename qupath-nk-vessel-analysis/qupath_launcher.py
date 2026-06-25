#!/usr/bin/env python3
"""
QuPath → Excel  Launcher
========================
GUI-Wrapper für qupath_to_excel.  Doppelklick auf die .exe öffnet ein
Fenster zum Ordner-Auswählen; das Ergebnis wird als
  <measurements_dir>/Analyse_export_<DATUM>/QuPath_Rohdaten.xlsx
gespeichert.

Bauen (auf Windows, im Ordner dieser Datei):
  pip install pyinstaller openpyxl
  pyinstaller --onefile --windowed --name QuPath_NK_Analysis qupath_launcher.py
→ dist/QuPath_NK_Analysis.exe
"""

import sys, os, csv, re, glob, threading, datetime, io
from collections import defaultdict
from statistics import mean

# ─── tkinter ──────────────────────────────────────────────────────────────────
import tkinter as tk
from tkinter import ttk, filedialog, scrolledtext, messagebox

# ─── openpyxl ─────────────────────────────────────────────────────────────────
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    Workbook = None

# ==============================================================================
#  qupath_to_excel  –  komplette Logik (keine externe Datei nötig)
# ==============================================================================

HB   = Font(bold=True)
GREY = PatternFill("solid", fgColor="DDDDDD")


def _find(measdir, stem):
    for ext in (".tsv", ".csv"):
        p = os.path.join(measdir, stem + ext)
        if os.path.exists(p):
            return p
    hits = glob.glob(os.path.join(measdir, stem + "*"))
    return hits[0] if hits else None


def _load(path):
    """Liest TSV/CSV; gibt (header, rows) zurück; Zahlen → float."""
    if path is None:
        return [], []
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096); f.seek(0)
        tc, sc, cc = sample.count("\t"), sample.count(";"), sample.count(",")
        delim = "\t" if (tc >= sc and tc >= cc) else (";" if sc >= cc else ",")
        rd = csv.reader(f, delimiter=delim)
        header = next(rd)
        rows = []
        for r in rd:
            d = {}
            for k, v in zip(header, r):
                if v is None or v == "":
                    d[k] = None
                else:
                    try:
                        d[k] = float(v.replace(",", ".")) if re.match(r"^-?\d", v) else v
                    except ValueError:
                        d[k] = v
            rows.append(d)
    return header, rows


def _parse_factors(img):
    s = str(img)
    return dict(
        animal=(re.search(r"_(\d{4,6})_", s) or [None, None])[1],
        sex=(re.search(r"_\d{4,6}_([fm])_", s) or [None, None])[1],
        cond="OP" if "gcOP" in s else ("N" if "gcN" in s else "?"),
        panel="IB4" if "IB4" in s else ("Syn" if "Syn" in s else "?"),
        sample=(re.search(r"_[fm]_([0-9]+\.[0-9]+)_", s) or [None, None])[1],
        is_control="control" in s.lower(),
    )


def _col_for(header, *cands):
    for c in cands:
        for h in header:
            if c.lower() in str(h).lower():
                return h
    return None


def _write_raw(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    for r in rows:
        ws.append([r.get(h) for h in header])
    ws.freeze_panes = "A2"


def build(measdir, out, log=print):
    """Erstellt QuPath_Rohdaten.xlsx aus den Messdateien in measdir."""
    if Workbook is None:
        raise RuntimeError("openpyxl ist nicht installiert.")

    log(f"Lese Dateien aus:\n  {measdir}\n")

    det_h, det  = _load(_find(measdir, "ALL_detections_NKcells"))
    non_h, non  = _load(_find(measdir, "ALL_detections_nonNK"))
    ves_h, ves  = _load(_find(measdir, "ALL_vessels_annotations"))
    sum_h, summ = _load(_find(measdir, "SUMMARY_per_image"))

    log(f"  NK-Zellen       : {len(det):>6}")
    log(f"  nonNK-Zellen    : {len(non):>6}")
    log(f"  Vessel-Annot.   : {len(ves):>6}")
    log(f"  Summary-Zeilen  : {len(summ):>6}\n")

    if not summ:
        raise FileNotFoundError(
            "SUMMARY_per_image nicht gefunden. "
            "Bitte den richtigen Measurements-Ordner wählen."
        )

    wb = Workbook(); wb.remove(wb.active)

    # ── Pivot_summary ──────────────────────────────────────────────────────────
    img_c = sum_h[0]
    ws = wb.create_sheet("Pivot_summary")
    factor_cols = ["animal", "sex", "cond", "panel", "sample"]
    head = [img_c] + factor_cols + sum_h[1:]
    ws.append(head)
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    numeric_cols = sum_h[1:]
    totals = defaultdict(list)
    for r in summ:
        f = _parse_factors(r[img_c])
        if f["is_control"]:
            continue
        ws.append([r[img_c]] + [f[k] for k in factor_cols] + [r.get(h) for h in sum_h[1:]])
        for h in numeric_cols:
            if isinstance(r.get(h), float):
                totals[h].append(r[h])
    grow = ["Gesamt/Mittel", "", "", "", "", ""]
    for h in numeric_cols:
        vals = totals.get(h, [])
        grow.append(
            round(sum(vals), 2) if ("count" in h or h.endswith("vessels") or h == "NK_count")
            else (round(mean(vals), 3) if vals else None)
        )
    ws.append(grow)
    for c in ws[ws.max_row]:
        c.font = HB
    ws.freeze_panes = "A2"
    log("  ✓ Pivot_summary")

    # ── Pivot_zones (Gefäß-Distanzzonen) ──────────────────────────────────────
    zone_pairs = [
        ("NK_count_zone_0_5um",    "pct_zone_0_5um",    "n 0-5µm",   "% 0-5µm"),
        ("NK_count_zone_5_10um",   "pct_zone_5_10um",   "n 5-10µm",  "% 5-10µm"),
        ("NK_count_zone_10_20um",  "pct_zone_10_20um",  "n 10-20µm", "% 10-20µm"),
        ("NK_count_zone_over20um", "pct_zone_over20um", "n >20µm",   "% >20µm"),
    ]
    zone_pairs = [(nc, pc, nl, pl) for nc, pc, nl, pl in zone_pairs if pc in sum_h]
    if zone_pairs:
        ws = wb.create_sheet("Pivot_zones")
        hdr = [img_c, "animal", "cond", "panel", "n_NK"]
        for _, _, nl, pl in zone_pairs:
            hdr += [nl, pl]
        hdr.append("Summe %")
        ws.append(hdr)
        for c in ws[1]:
            c.font = HB; c.fill = GREY
        ptot = defaultdict(list)
        for r in summ:
            f = _parse_factors(r[img_c])
            if f["is_control"]:
                continue
            row = [r[img_c], f["animal"], f["cond"], f["panel"], r.get("NK_count")]
            psum = 0.0
            for nc, pc, _, _ in zone_pairs:
                nv = r.get(nc); pv = r.get(pc)
                row += [int(nv) if isinstance(nv, float) else nv,
                        round(pv, 2) if isinstance(pv, float) else None]
                if isinstance(pv, float):
                    psum += pv; ptot[pc].append(pv)
            row.append(round(psum, 1))
            ws.append(row)
        grow = ["Mittel über Bilder", "", "", "", ""]
        for _, pc, _, _ in zone_pairs:
            vv = ptot.get(pc, [])
            grow += [None, round(mean(vv), 2) if vv else None]
        grow.append(round(sum(g for g in grow[5:] if isinstance(g, (int, float))), 1))
        ws.append(grow)
        for c in ws[ws.max_row]:
            c.font = HB
        ws.freeze_panes = "A2"
        log("  ✓ Pivot_zones")

    # ── Pivot_zones_syn (Synaptophysin-Distanzzonen) ──────────────────────────
    syn_src = "Dist synaptophysin um"
    if syn_src in det_h:
        img_dd = det_h[0]

        def _syn_zone(d):
            return 0 if d <= 5.0 else (1 if d <= 10.0 else (2 if d <= 20.0 else 3))

        zlabels = [("n 0-5µm", "% 0-5µm"), ("n 5-10µm", "% 5-10µm"),
                   ("n 10-20µm", "% 10-20µm"), ("n >20µm", "% >20µm")]
        zc = defaultdict(lambda: [0, 0, 0, 0])
        for r in det:
            v = r.get(syn_src)
            if isinstance(v, float):
                zc[r[img_dd]][_syn_zone(v)] += 1
        if zc:
            ws = wb.create_sheet("Pivot_zones_syn")
            hdr = [img_dd, "animal", "cond", "panel", "n_NK"]
            for nl, pl in zlabels:
                hdr += [nl, pl]
            hdr.append("Summe %")
            ws.append(hdr)
            for c in ws[1]:
                c.font = HB; c.fill = GREY
            ptot = defaultdict(list)
            for im in sorted(zc):
                f = _parse_factors(im)
                if f["is_control"]:
                    continue
                counts = zc[im]; tot = sum(counts)
                row = [im, f["animal"], f["cond"], f["panel"], tot]
                psum = 0.0
                for k, (nl, pl) in enumerate(zlabels):
                    pct = (counts[k] * 100.0 / tot) if tot else None
                    row += [counts[k], round(pct, 2) if pct is not None else None]
                    if pct is not None:
                        psum += pct; ptot[k].append(pct)
                row.append(round(psum, 1))
                ws.append(row)
            grow = ["Mittel über Bilder", "", "", "", ""]
            for k in range(4):
                vv = ptot.get(k, [])
                grow += [None, round(mean(vv), 2) if vv else None]
            grow.append(round(sum(g for g in grow[5:] if isinstance(g, (int, float))), 1))
            ws.append(grow)
            for c in ws[ws.max_row]:
                c.font = HB
            ws.freeze_panes = "A2"
            log("  ✓ Pivot_zones_syn")

    # ── Pivot_NKcells (Morphologie + Distanzen je Bild) ───────────────────────
    img_d = det_h[0]
    pick = [
        ("Cell: Area",          _col_for(det_h, "Cell: Area")),
        ("Cell: Perimeter",     "Cell: Perimeter"),
        ("Cell: Circularity",   "Cell: Circularity"),
        ("Cell: Max caliper",   "Cell: Max caliper"),
        ("Cell: Min caliper",   "Cell: Min caliper"),
        ("Cell: NKp46 mean",    "Cell: NKp46 mean"),
        ("Cell: NKp46 std dev", "Cell: NKp46 std dev"),
        ("Cell: NKp46 max",     "Cell: NKp46 max"),
        ("Dist small vessel um",   "Dist small vessel um"),
        ("Dist large vessel um",   "Dist large vessel um"),
        ("Dist synaptophysin um",  "Dist synaptophysin um"),
        ("Dist nearest vessel um", "Dist nearest vessel um"),
        ("Perivascular (0/1)",     "Perivascular (0/1)"),
    ]
    pick = [(lbl, src) for lbl, src in pick if src and src in det_h]
    agg = defaultdict(lambda: defaultdict(list)); cnt = defaultdict(int)
    for r in det:
        im = r[img_d]; cnt[im] += 1
        for lbl, src in pick:
            v = r.get(src)
            if isinstance(v, float):
                agg[im][lbl].append(v)
    ws = wb.create_sheet("Pivot_NKcells")
    ws.append([img_d, "animal", "cond", "panel", "n_NK"] + [lbl for lbl, _ in pick])
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    for im in sorted(cnt):
        f = _parse_factors(im)
        row = [im, f["animal"], f["cond"], f["panel"], cnt[im]]
        for lbl, _ in pick:
            vals = agg[im][lbl]
            row.append(round(mean(vals), 3) if vals else None)
        ws.append(row)
    ws.freeze_panes = "A2"
    log("  ✓ Pivot_NKcells (inkl. Morphologie)")

    # ── Pivot_nonNKcells ──────────────────────────────────────────────────────
    if non:
        img_n = non_h[0]
        dist_pick = [(lbl, src) for lbl, src in [
            ("Dist small vessel um",   "Dist small vessel um"),
            ("Dist large vessel um",   "Dist large vessel um"),
            ("Dist nearest vessel um", "Dist nearest vessel um"),
            ("Dist synaptophysin um",  "Dist synaptophysin um"),
            ("Perivascular (0/1)",     "Perivascular (0/1)"),
        ] if src in non_h]
        nagg = defaultdict(lambda: defaultdict(list)); ncnt = defaultdict(int)
        for r in non:
            im = r[img_n]; ncnt[im] += 1
            for lbl, src in dist_pick:
                v = r.get(src)
                if isinstance(v, float):
                    nagg[im][lbl].append(v)
        ws = wb.create_sheet("Pivot_nonNKcells")
        ws.append([img_n, "animal", "cond", "panel", "n_nonNK"] + [lbl for lbl, _ in dist_pick])
        for c in ws[1]:
            c.font = HB; c.fill = GREY
        for im in sorted(ncnt):
            f = _parse_factors(im)
            row = [im, f["animal"], f["cond"], f["panel"], ncnt[im]]
            for lbl, _ in dist_pick:
                vals = nagg[im][lbl]
                row.append(round(mean(vals), 3) if vals else None)
            ws.append(row)
        ws.freeze_panes = "A2"
        log("  ✓ Pivot_nonNKcells")

    # ── Pivot_vessels ─────────────────────────────────────────────────────────
    img_v = ves_h[0]
    vcols = [c for c in ["Num Detections", "Num NK-cell",
                          "NK within 20um count", "NK within 5um count"] if c in ves_h]
    vagg = defaultdict(lambda: defaultdict(float)); vcnt = defaultdict(int)
    for r in ves:
        im = r[img_v]; vcnt[im] += 1
        for c in vcols:
            v = r.get(c)
            if isinstance(v, float):
                vagg[im][c] += v
    ws = wb.create_sheet("Pivot_vessels")
    ws.append([img_v, "animal", "cond", "panel", "n_annot"] + [f"Summe {c}" for c in vcols])
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    for im in sorted(vcnt):
        f = _parse_factors(im)
        ws.append([im, f["animal"], f["cond"], f["panel"], vcnt[im]] +
                  [round(vagg[im][c], 1) for c in vcols])
    ws.freeze_panes = "A2"
    log("  ✓ Pivot_vessels")

    # ── Pivot_struct_area ─────────────────────────────────────────────────────
    area_col = _col_for(ves_h, "Area µm^2", "Area um^2", "Area")
    cls_col  = next((h for h in ves_h if str(h).strip().lower() == "classification"), None)
    if area_col and cls_col:
        st = defaultdict(lambda: {"tissue_um2": 0.0, "syn": [], "small": [], "large": []})
        for r in ves:
            cls = str(r.get(cls_col) or "").strip()
            a   = r.get(area_col)
            im  = r[img_v]
            if not isinstance(a, float):
                continue
            if   cls == "Tissue":        st[im]["tissue_um2"] += a
            elif cls == "Synaptophysin": st[im]["syn"].append(a)
            elif cls == "small vessel":  st[im]["small"].append(a)
            elif cls == "large vessel":  st[im]["large"].append(a)
        ws = wb.create_sheet("Pivot_struct_area")
        ws.append([img_v, "animal", "cond", "panel",
                   "tissue_area_mm2",
                   "n_synaptophysin", "syn_density_per_mm2",
                   "syn_area_mean_um2", "syn_area_fraction_pct",
                   "n_small_vessel", "small_vessel_area_mean_um2",
                   "n_large_vessel", "large_vessel_area_mean_um2"])
        for c in ws[1]:
            c.font = HB; c.fill = GREY
        for im in sorted(st):
            f = _parse_factors(im)
            if f["is_control"]:
                continue
            s = st[im]; tmm2 = s["tissue_um2"] / 1e6 if s["tissue_um2"] else None
            syn, sm, lg = s["syn"], s["small"], s["large"]
            ws.append([
                im, f["animal"], f["cond"], f["panel"],
                round(tmm2, 4) if tmm2 else None,
                len(syn),
                round(len(syn) / tmm2, 2) if (syn and tmm2) else None,
                round(mean(syn), 2) if syn else None,
                round(sum(syn) * 100.0 / s["tissue_um2"], 4) if (syn and s["tissue_um2"]) else None,
                len(sm),
                round(mean(sm), 2) if sm else None,
                len(lg),
                round(mean(lg), 2) if lg else None,
            ])
        ws.freeze_panes = "A2"
        log("  ✓ Pivot_struct_area")

    # ── PerAnimal_condition ───────────────────────────────────────────────────
    metrics = [c for c in [
        "NK_density_per_mm2",
        "small_vessel_density_per_mm2", "large_vessel_density_per_mm2", "vessel_density_per_mm2",
        "perivascular_pct", "vessel_area_fraction_pct",
        "mean_dist_small_um", "mean_dist_large_um", "mean_dist_syn_um",
        "NK_count_zone_0_5um",    "pct_zone_0_5um",
        "NK_count_zone_5_10um",   "pct_zone_5_10um",
        "NK_count_zone_10_20um",  "pct_zone_10_20um",
        "NK_count_zone_over20um", "pct_zone_over20um",
    ] if c in sum_h]
    pa = defaultdict(list)
    for r in summ:
        f = _parse_factors(r[img_c])
        if f["is_control"]:
            continue
        for m in metrics:
            if isinstance(r.get(m), float):
                pa[(f["animal"], f["cond"], m)].append(r[m])
    ws = wb.create_sheet("PerAnimal_condition")
    ws.append(["animal", "cond"] + metrics)
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    animals = sorted({
        _parse_factors(r[img_c])["animal"]
        for r in summ
        if not _parse_factors(r[img_c])["is_control"]
    })
    for a in animals:
        for cond in ["OP", "N"]:
            row = [a, cond]
            for m in metrics:
                vals = pa.get((a, cond, m), [])
                row.append(round(mean(vals), 2) if vals else None)
            ws.append(row)
    ws.freeze_panes = "A2"
    log("  ✓ PerAnimal_condition")

    # ── Grunddaten (Rohkopien) ────────────────────────────────────────────────
    _write_raw(wb, "Grunddaten_summary", sum_h, summ)
    _write_raw(wb, "Grunddaten_vessels", ves_h, ves)
    _write_raw(wb, "Grunddaten_NKcells", det_h, det)
    if non:
        _write_raw(wb, "Grunddaten_nonNK", non_h, non)
    log("  ✓ Grunddaten_*\n")

    wb.save(out)
    log(f"Gespeichert ({len(wb.sheetnames)} Sheets):\n  {out}\n")


# ==============================================================================
#  GUI
# ==============================================================================

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("QuPath → Excel  (NK-Analyse)")
        self.resizable(True, True)
        self.minsize(620, 440)
        self._build_ui()
        self._center()

    def _center(self):
        self.update_idletasks()
        w, h = 680, 500
        sw = self.winfo_screenwidth(); sh = self.winfo_screenheight()
        self.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")

    def _build_ui(self):
        pad = dict(padx=12, pady=6)

        # ── Ordner-Zeile ──────────────────────────────────────────────────────
        frm = ttk.LabelFrame(self, text="Measurements-Ordner (QuPath-Export)")
        frm.pack(fill="x", **pad)

        self._dir_var = tk.StringVar()
        ttk.Entry(frm, textvariable=self._dir_var, width=60).pack(
            side="left", fill="x", expand=True, padx=(6, 4), pady=6)
        ttk.Button(frm, text="Durchsuchen …", command=self._pick_dir).pack(
            side="left", padx=(0, 6), pady=6)

        # ── Ausgabe-Zeile ─────────────────────────────────────────────────────
        frm2 = ttk.LabelFrame(self, text="Ausgabe-Datei  (leer lassen = automatisch)")
        frm2.pack(fill="x", **pad)

        self._out_var = tk.StringVar()
        ttk.Entry(frm2, textvariable=self._out_var, width=60).pack(
            side="left", fill="x", expand=True, padx=(6, 4), pady=6)
        ttk.Button(frm2, text="Speichern unter …", command=self._pick_out).pack(
            side="left", padx=(0, 6), pady=6)

        # ── Start-Button ──────────────────────────────────────────────────────
        self._btn = ttk.Button(self, text="▶  Analyse starten",
                               command=self._run, style="Accent.TButton")
        self._btn.pack(pady=(4, 2))

        # ── Log-Feld ──────────────────────────────────────────────────────────
        frm3 = ttk.LabelFrame(self, text="Protokoll")
        frm3.pack(fill="both", expand=True, padx=12, pady=(0, 4))
        self._log = scrolledtext.ScrolledText(frm3, height=12, state="disabled",
                                              font=("Consolas", 9), wrap="word")
        self._log.pack(fill="both", expand=True, padx=4, pady=4)

        # ── Statusleiste ──────────────────────────────────────────────────────
        self._status = tk.StringVar(value="Bereit.")
        ttk.Label(self, textvariable=self._status, anchor="w",
                  relief="sunken").pack(fill="x", side="bottom")

    # ── Dateiauswahl ──────────────────────────────────────────────────────────
    def _pick_dir(self):
        d = filedialog.askdirectory(title="Measurements-Ordner wählen")
        if d:
            self._dir_var.set(d)

    def _pick_out(self):
        f = filedialog.asksaveasfilename(
            title="Ausgabe-Datei wählen",
            defaultextension=".xlsx",
            filetypes=[("Excel-Arbeitsmappe", "*.xlsx")],
        )
        if f:
            self._out_var.set(f)

    # ── Analyse ───────────────────────────────────────────────────────────────
    def _run(self):
        measdir = self._dir_var.get().strip()
        if not measdir or not os.path.isdir(measdir):
            messagebox.showerror("Fehler", "Bitte einen gültigen Ordner auswählen.")
            return
        if Workbook is None:
            messagebox.showerror(
                "Fehler",
                "openpyxl ist nicht installiert.\n"
                "Bitte in der Eingabeaufforderung ausführen:\n"
                "  pip install openpyxl",
            )
            return

        out = self._out_var.get().strip()
        if not out:
            stamp = datetime.date.today().strftime("%Y-%m-%d")
            export_dir = os.path.join(measdir, f"Analyse_export_{stamp}")
            os.makedirs(export_dir, exist_ok=True)
            out = os.path.join(export_dir, "QuPath_Rohdaten.xlsx")

        self._log_clear()
        self._btn.config(state="disabled")
        self._status.set("Läuft …")

        def worker():
            try:
                build(measdir, out, log=self._log_append)
                self.after(0, self._done, out)
            except Exception as exc:
                self.after(0, self._error, str(exc))

        threading.Thread(target=worker, daemon=True).start()

    def _done(self, out):
        self._btn.config(state="normal")
        self._status.set(f"Fertig → {out}")
        messagebox.showinfo(
            "Fertig",
            f"Excel-Datei gespeichert:\n\n{out}",
        )

    def _error(self, msg):
        self._btn.config(state="normal")
        self._status.set("Fehler!")
        self._log_append(f"\n[FEHLER] {msg}\n")
        messagebox.showerror("Fehler", msg)

    # ── Log-Hilfen ────────────────────────────────────────────────────────────
    def _log_append(self, text):
        def _do():
            self._log.config(state="normal")
            self._log.insert("end", text + "\n")
            self._log.see("end")
            self._log.config(state="disabled")
        self.after(0, _do)

    def _log_clear(self):
        self._log.config(state="normal")
        self._log.delete("1.0", "end")
        self._log.config(state="disabled")


# ==============================================================================
#  Einstiegspunkt
# ==============================================================================

if __name__ == "__main__":
    app = App()
    app.mainloop()

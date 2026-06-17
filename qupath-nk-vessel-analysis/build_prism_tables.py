import openpyxl, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict
from statistics import mean, median

SRC="/root/.claude/uploads/156810d6-bed2-5a8a-a441-9106e2f41e94/0b8e9e48-Export_neu.xlsx"
OUT="/home/user/.github/qupath-nk-vessel-analysis/Prism_ready_analysis.xlsx"
wb=openpyxl.load_workbook(SRC, data_only=True, read_only=True)

def parse(img):
    s=str(img)
    animal=(re.search(r'_(\d{4,6})_',s) or [None,None])[1]
    sex=(re.search(r'_\d{4,6}_([fm])_',s) or [None,None])[1]
    cond='OP' if 'gcOP' in s else ('N' if 'gcN' in s else '?')
    panel='IB4' if 'IB4' in s else ('Syn' if 'Syn' in s else '?')
    samp=(re.search(r'_[fm]_([0-9]+\.[0-9]+)_',s) or [None,None])[1]
    ctrl='control' in s.lower()
    return dict(animal=animal,sex=sex,cond=cond,panel=panel,sample=samp,ctrl=ctrl)

# ---- read SUMMARY_per_image ----
ws=wb["SUMMARY_per_image"]; rows=list(ws.iter_rows(values_only=True)); H=rows[0]
def col(name): return H.index(name)
recs=[]
for r in rows[1:]:
    if r[0] is None: continue
    d=parse(r[0]); 
    if d['ctrl']: continue
    d['image']=str(r[0])
    for nm in ["NK_count","NK_density_per_mm2","vessel_area_fraction_pct","perivascular_pct",
               "mean_dist_small_um","mean_dist_large_um","mean_dist_syn_um",
               "median_dist_small_um","median_dist_large_um","median_dist_syn_um",
               "small_vessels","large_vessels"]:
        v=r[col(nm)]
        d[nm]=v
    d['animal_cond']=f"{d['animal']}_{d['cond']}"
    recs.append(d)

# ---- per-cell activation-by-location (IB4 only) ----
wsd=wb["ALL_detections_NKcells"]; itd=wsd.iter_rows(values_only=True); Hd=next(itd)
def dc(n): return Hd.index(n)
img_i=dc('Image'); peri_i=dc('Perivascular (0/1)'); nk_i=dc('Cell: NKp46 mean'); area_i=dc('Cell: Area')
ds_i=dc('Dist small vessel um'); dl_i=dc('Dist large vessel um')
peri_nk=defaultdict(list); nonperi_nk=defaultdict(list)
peri_area=defaultdict(list); nonperi_area=defaultdict(list)
for r in itd:
    img=str(r[img_i]); p=r[peri_i]
    if p is None: continue
    nkv=r[nk_i]; ar=r[area_i]
    if p==1:
        if nkv is not None: peri_nk[img].append(nkv)
        if ar is not None: peri_area[img].append(ar)
    else:
        if nkv is not None: nonperi_nk[img].append(nkv)
        if ar is not None: nonperi_area[img].append(ar)

# ---- per-vessel NK-within by caliber ----
wsv=wb["ALL_vessels_annotations"]; itv=wsv.iter_rows(values_only=True); Hv=next(itv)
vi=Hv.index('Image'); vci=Hv.index('Classification'); vw20=Hv.index('NK within 20um count')
cal=defaultdict(lambda: defaultdict(list))  # image -> class -> [counts]
for r in itv:
    cls=r[vci]; 
    if cls not in ('small vessel','large vessel'): continue
    w=r[vw20]
    if w is None: continue
    cal[str(r[vi])][cls].append(w)

# =================== build workbook ===================
out=Workbook()
H1=Font(bold=True,size=13); HB=Font(bold=True); GREY=PatternFill("solid",fgColor="DDDDDD")
def style_header(ws,row=1):
    for c in ws[row]:
        c.font=HB; c.fill=GREY; c.alignment=Alignment(horizontal="center")

# ---- README ----
ws0=out.active; ws0.title="README_Analyseplan"
readme=[
("NK-Zell / Gefäß / Nerv – Analyseplan & Prism-Tabellen",H1),
("",None),
("Design (aus Dateinamen dekodiert):",HB),
("  Tiere (biol. Replikat): 2  — 86426 (m), 89807 (f).  Genotyp: nur WT.",None),
("  Bedingung: OP (operiert) vs N (naiv/kontralateral) — beide Tiere haben beide Seiten.",None),
("  Panel: IB4 (Gefäße: small/large) vs Synaptophysin (Nerventerminals/NMJ) – SEPARATE Schnitte.",None),
("  ~3 Bilder je Schnitt; Kontroll-/Negativbilder (NK=0) ausgeschlossen.",None),
("",None),
("*** WICHTIGSTER STAT-HINWEIS ***",HB),
("  Biologisches Replikat = TIER, nicht Bild und nicht Zelle. Hier n=2 Tiere.",None),
("  => Mit n=2 sind p-Werte über Tiere NICHT belastbar (rein deskriptiv / Pilot).",None),
("  => Zellen/Bilder als 'n' zu verwenden ist Pseudoreplikation (Lazic 2010; Aarts 2014).",None),
("  Empfehlung: Pilot beschreiben (Effektgrößen + Einzeltierpunkte zeigen);",None),
("  für Inferenz n>=4-5 Tiere/Gruppe sammeln und gemischte Modelle (Bild in Tier genestet) nutzen.",None),
("",None),
("Vorgeschlagene Analysen (Sheet -> Prism-Test):",HB),
("  A1  NK-Dichte OP vs N (IB4 & Syn getrennt) -> Mann-Whitney / ungepaart; Tier-Mittel paarweise.",None),
("  A2  Perivaskulärer Anteil % (NK <20µm an Gefäß) OP vs N -> Mann-Whitney.",None),
("  A3  Distanz NK->small vs ->large Gefäß (gepaart je IB4-Bild) -> Wilcoxon paired.",None),
("  A4  Distanz NK->Synaptophysin OP vs N -> Mann-Whitney.",None),
("  A5  NK im 20µm-Umkreis je Gefäß: small vs large Kaliber -> Mann-Whitney/paired je Bild.",None),
("  A6  Aktivierung nach Lage: NKp46-Intensität & Zellfläche perivaskulär vs nicht (gepaart je Bild) -> Wilcoxon paired.",None),
("",None),
("WICHTIG – Anreicherungstest fehlt noch (empfohlen!):",HB),
("  Um zu zeigen, dass NK-Zellen NÄHER an Gefäßen/Nerven liegen als zufällig,",None),
("  vergleiche NK- gegen Nicht-NK-Zell-Distanzen in DENSELBEN Schnitten",None),
("  (oder gegen permutierte Zufallspositionen / Ripley-K).",None),
("  Der aktuelle Export enthält nur NK-Distanzen. Ich kann das QuPath-Skript erweitern,",None),
("  sodass es auch die Nicht-NK-Hintergrunddistanzen exportiert.",None),
("",None),
("Referenzen (real, bitte domänenspezifische selbst gegenprüfen):",HB),
("  Spatial point patterns: Ripley BD 1977 JRSS-B; Baddeley/Rubak/Turner 2015 (spatstat-Buch); Diggle 2013.",None),
("  Nachbarschafts-/Anreicherungstests in Gewebe-Imaging: Schapiro 2017 (histoCAT) Nat Methods;",None),
("    Palla 2022 (squidpy) Nat Methods; Windhager 2023 (imcRtools/steinbock) Nat Protocols.",None),
("  QuPath-Methodik: Bankhead 2017 Sci Rep.",None),
("  Pseudoreplikation/genestete Daten: Lazic 2010 BMC Neurosci; Aarts 2014 Nat Neurosci.",None),
("  Immunzellen im Muskel (Kontext): Tidball 2017 Nat Rev Immunol.",None),
]
for i,(t,f) in enumerate(readme,1):
    c=ws0.cell(row=i,column=1,value=t)
    if f: c.font=f
ws0.column_dimensions['A'].width=110

def block_sheet(title, group_cols, data_by_group, note, decimals=2):
    ws=out.create_sheet(title)
    ws.cell(1,1,note).font=HB
    for j,g in enumerate(group_cols,1):
        ws.cell(3,j,g).font=HB; ws.cell(3,j).fill=GREY
    maxn=max((len(data_by_group[g]) for g in group_cols), default=0)
    for i in range(maxn):
        for j,g in enumerate(group_cols,1):
            vals=data_by_group[g]
            if i<len(vals) and vals[i] is not None:
                ws.cell(4+i,j, round(float(vals[i]),decimals))
    for j in range(1,len(group_cols)+1):
        ws.column_dimensions[chr(64+j)].width=16
    return ws

# ---- A1 density ----
for panel in ["IB4","Syn"]:
    g={"OP":[],"N":[]}
    for d in recs:
        if d['panel']==panel and d['NK_density_per_mm2'] is not None:
            g[d['cond']].append(d['NK_density_per_mm2'])
    block_sheet(f"A1_NKdensity_{panel}",["OP","N"],g,
        f"NK-Dichte (Zellen/mm²), Panel {panel}, je Bild. Prism: Mann-Whitney (deskriptiv, n=2 Tiere!).")

# ---- A2 perivascular pct (IB4) ----
g={"OP":[],"N":[]}
for d in recs:
    if d['panel']=='IB4' and d['perivascular_pct'] is not None:
        g[d['cond']].append(d['perivascular_pct'])
block_sheet("A2_Perivascular_pct",["OP","N"],g,
    "Perivaskulärer Anteil [% NK <20µm am Gefäß], je IB4-Bild. Prism: Mann-Whitney.")

# ---- A3 small vs large (paired per IB4 image) ----
ws=out.create_sheet("A3_Dist_small_vs_large")
ws.cell(1,1,"Mittl. Distanz NK->Gefäß je IB4-Bild, gepaart small vs large. Prism: Wilcoxon paired.").font=HB
for j,h in enumerate(["Bild","cond","small_um","large_um"],1):
    ws.cell(3,j,h).font=HB; ws.cell(3,j).fill=GREY
ri=4
for d in recs:
    if d['panel']=='IB4' and d['mean_dist_small_um'] is not None and d['mean_dist_large_um'] is not None:
        ws.cell(ri,1,d['image'][:40]); ws.cell(ri,2,d['cond'])
        ws.cell(ri,3,round(d['mean_dist_small_um'],2)); ws.cell(ri,4,round(d['mean_dist_large_um'],2)); ri+=1
ws.column_dimensions['A'].width=42

# ---- A4 dist to syn OP vs N ----
g={"OP":[],"N":[]}
for d in recs:
    if d['panel']=='Syn' and d['mean_dist_syn_um'] is not None:
        g[d['cond']].append(d['mean_dist_syn_um'])
block_sheet("A4_Dist_to_Syn",["OP","N"],g,
    "Mittl. Distanz NK->Synaptophysin [µm], je Syn-Bild. Prism: Mann-Whitney.")

# ---- A5 caliber NK within 20 (mean per image small vs large) ----
ws=out.create_sheet("A5_Caliber_NKwithin20")
ws.cell(1,1,"Mittl. NK-Zahl im 20µm-Umkreis je Gefäß; gepaart small vs large je IB4-Bild. Prism: Wilcoxon paired.").font=HB
for j,h in enumerate(["Bild","small_meanNK","large_meanNK"],1):
    ws.cell(3,j,h).font=HB; ws.cell(3,j).fill=GREY
ri=4
for img,cc in cal.items():
    if 'small vessel' in cc and 'large vessel' in cc:
        ws.cell(ri,1,img[:40])
        ws.cell(ri,2,round(mean(cc['small vessel']),3))
        ws.cell(ri,3,round(mean(cc['large vessel']),3)); ri+=1
ws.column_dimensions['A'].width=42

# ---- A6 activation by location (paired per IB4 image) ----
ws=out.create_sheet("A6_Activation_by_location")
ws.cell(1,1,"Pro IB4-Bild: Mittelwert NKp46 (Cell) & Zellfläche, perivaskulär vs nicht. Prism: Wilcoxon paired.").font=HB
hdr=["Bild","cond","NKp46_peri","NKp46_nonPeri","Area_peri","Area_nonPeri"]
for j,h in enumerate(hdr,1):
    ws.cell(3,j,h).font=HB; ws.cell(3,j).fill=GREY
ri=4
imgset={d['image']:d for d in recs if d['panel']=='IB4'}
for img,d in imgset.items():
    if img in peri_nk and img in nonperi_nk and peri_nk[img] and nonperi_nk[img]:
        ws.cell(ri,1,img[:40]); ws.cell(ri,2,d['cond'])
        ws.cell(ri,3,round(mean(peri_nk[img]),1)); ws.cell(ri,4,round(mean(nonperi_nk[img]),1))
        ws.cell(ri,5,round(mean(peri_area[img]),2)); ws.cell(ri,6,round(mean(nonperi_area[img]),2)); ri+=1
ws.column_dimensions['A'].width=42

# ---- Animal-level means (proper replicate) ----
ws=out.create_sheet("AnimalLevel_means")
ws.cell(1,1,"Tier-Mittelwerte (richtiges biol. Replikat, n=2). Für gepaarte OP-vs-N-Darstellung in Prism.").font=HB
metrics=["NK_density_per_mm2","perivascular_pct","mean_dist_small_um","mean_dist_large_um","mean_dist_syn_um"]
agg=defaultdict(list)
for d in recs:
    for m in metrics:
        if d[m] is not None: agg[(d['animal'],d['cond'],m)].append(d[m])
hdr=["animal","cond"]+metrics
for j,h in enumerate(hdr,1): ws.cell(3,j,h).font=HB; ws.cell(3,j).fill=GREY
ri=4
for animal in sorted(set(d['animal'] for d in recs)):
    for cond in ["OP","N"]:
        ws.cell(ri,1,animal); ws.cell(ri,2,cond)
        for j,m in enumerate(metrics,3):
            vals=agg.get((animal,cond,m))
            if vals: ws.cell(ri,j,round(mean(vals),2))
        ri+=1
for j in range(1,len(hdr)+1): ws.column_dimensions[chr(64+j)].width=16

# ---- tidy per-image ----
ws=out.create_sheet("PerImage_tidy")
cols=["image","animal","sex","cond","panel","sample","NK_count","NK_density_per_mm2",
      "vessel_area_fraction_pct","perivascular_pct","mean_dist_small_um","mean_dist_large_um","mean_dist_syn_um"]
for j,c in enumerate(cols,1): ws.cell(1,j,c).font=HB; ws.cell(1,j).fill=GREY
for i,d in enumerate(recs,2):
    for j,c in enumerate(cols,1):
        v=d.get(c); ws.cell(i,j, round(v,3) if isinstance(v,float) else v)
ws.column_dimensions['A'].width=42

out.save(OUT)
print("WROTE", OUT)
print("recs(images, no controls):",len(recs))

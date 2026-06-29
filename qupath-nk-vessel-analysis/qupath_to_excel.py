#!/usr/bin/env python3
"""
QuPath-Exporte  ->  fertige Excel-Arbeitsmappe (Grunddatensatz + Pivot-Tabellen)
================================================================================
Liest die vom Groovy-Skript erzeugten Tabellen aus <Projekt>/measurements/ und
baut EINE .xlsx-Mappe mit:
  - Grunddaten_*            : Rohdaten (Detections / Vessels / Summary)
  - Pivot_summary           : je Bild Kennzahlen (+ Faktoren animal/cond/panel) + Gesamt
  - Pivot_NKcells           : je Bild Mittelwerte (Morphologie, NKp46, Distanzen)
  - Pivot_vessels           : je Bild Summen (Detections, NK, NK-within)
  - Pivot_struct_area       : je Bild Dichte & Fläche (Synaptophysin, small/large vessel)
  - PerAnimal_condition     : Tier x Bedingung Mittelwerte (richtiges biol. Replikat)

Eingaben (TSV oder CSV; Trennzeichen wird erkannt):
  ALL_detections_NKcells.(tsv|csv)
  ALL_vessels_annotations.(tsv|csv)
  SUMMARY_per_image.(tsv|csv)

Aufruf:
  python3 qupath_to_excel.py  <measurements_dir>  [output.xlsx]
"""
import sys, os, csv, re, glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from collections import defaultdict
from statistics import mean

HB = Font(bold=True); GREY = PatternFill("solid", fgColor="DDDDDD")

# ---------------------------- Einlesen ----------------------------
def find(measdir, stem):
    for ext in (".tsv", ".csv"):
        p = os.path.join(measdir, stem + ext)
        if os.path.exists(p):
            return p
    hits = glob.glob(os.path.join(measdir, stem + "*"))
    return hits[0] if hits else None

def load(path):
    """Liest TSV/CSV, gibt (header:list, rows:list[dict]) zurück; Zahlen -> float."""
    if path is None:
        return [], []
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096); f.seek(0)
        delim = "\t" if sample.count("\t") >= sample.count(";") and sample.count("\t") >= sample.count(",") \
                else (";" if sample.count(";") >= sample.count(",") else ",")
        rd = csv.reader(f, delimiter=delim)
        header = next(rd)
        rows = []
        for r in rd:
            d = {}
            for k, v in zip(header, r):
                if v is None or v == "":
                    d[k] = None
                else:
                    try:
                        d[k] = float(v.replace(",", ".")) if re.match(r"^-?\d", v) else v
                    except ValueError:
                        d[k] = v
            rows.append(d)
    return header, rows

def parse_factors(img):
    s = str(img)
    return dict(
        animal=(re.search(r"_(\d{4,6})_", s) or [None, None])[1],
        sex=(re.search(r"_\d{4,6}_([fm])_", s) or [None, None])[1],
        cond="OP" if "gcOP" in s else ("N" if "gcN" in s else "?"),
        panel="IB4" if "IB4" in s else ("Syn" if "Syn" in s else "?"),
        sample=(re.search(r"_[fm]_([0-9]+\.[0-9]+)_", s) or [None, None])[1],
        is_control="control" in s.lower(),
    )

# ---------------------------- Helpers ----------------------------
def col_for(header, *cands):
    """Erste Spalte, deren Name eine der Teilzeichenketten enthält (robust ggü. µm/Âµm)."""
    for c in cands:
        for h in header:
            if c.lower() in str(h).lower():
                return h
    return None

def write_raw(wb, title, header, rows, maxcols=None):
    ws = wb.create_sheet(title)
    cols = header[:maxcols] if maxcols else header
    ws.append(cols)
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    for r in rows:
        ws.append([r.get(h) for h in cols])
    ws.freeze_panes = "A2"
    return ws

# ---------------------------- Build ----------------------------
def build(measdir, out):
    det_h, det  = load(find(measdir, "ALL_detections_NKcells"))
    non_h, non  = load(find(measdir, "ALL_detections_nonNK"))
    ves_h, ves  = load(find(measdir, "ALL_vessels_annotations"))
    sum_h, summ = load(find(measdir, "SUMMARY_per_image"))
    print(f"  NK={len(det)}  nonNK={len(non)}  vessels={len(ves)}  summary={len(summ)}")

    wb = Workbook(); wb.remove(wb.active)

    # ── Struktur-Metriken vorab berechnen (aus ALL_vessels_annotations) ──────
    # Wird in Pivot_summary, PerAnimal_condition und Pivot_struct_area genutzt.
    area_col = col_for(ves_h, "Area µm^2", "Area um^2", "Area")
    cls_col  = next((h for h in ves_h if str(h).strip().lower() == "classification"), None)
    img_v    = ves_h[0] if ves_h else None

    # struct_img[image] -> dict mit 5 gemittelten Area-Kennzahlen je Bild
    struct_img = {}
    # struct_raw[image] -> Rohdaten für Pivot_struct_area (Counts + Listen)
    struct_raw = {}
    STRUCT_COLS = ["syn_density_per_mm2", "syn_area_mean_um2", "syn_area_fraction_pct",
                   "small_vessel_area_mean_um2", "large_vessel_area_mean_um2"]
    if area_col and cls_col and img_v:
        _raw = defaultdict(lambda: {"tissue_um2": 0.0, "syn": [], "small": [], "large": []})
        for r in ves:
            _c = str(r.get(cls_col) or "").strip()
            _a = r.get(area_col)
            _i = r[img_v]
            if not isinstance(_a, float):
                continue
            if   _c == "Tissue":        _raw[_i]["tissue_um2"] += _a
            elif _c == "Synaptophysin": _raw[_i]["syn"].append(_a)
            elif _c == "small vessel":  _raw[_i]["small"].append(_a)
            elif _c == "large vessel":  _raw[_i]["large"].append(_a)
        for _i, _s in _raw.items():
            _t = _s["tissue_um2"] / 1e6 if _s["tissue_um2"] else None
            struct_img[_i] = {
                "syn_density_per_mm2":       round(len(_s["syn"]) / _t, 2)           if (_s["syn"] and _t)           else None,
                "syn_area_mean_um2":          round(mean(_s["syn"]), 2)               if _s["syn"]                    else None,
                "syn_area_fraction_pct":      round(sum(_s["syn"]) * 100.0 / _s["tissue_um2"], 4) if (_s["syn"] and _s["tissue_um2"]) else None,
                "small_vessel_area_mean_um2": round(mean(_s["small"]), 2)             if _s["small"]                  else None,
                "large_vessel_area_mean_um2": round(mean(_s["large"]), 2)             if _s["large"]                  else None,
            }
            struct_raw[_i] = _s   # für Pivot_struct_area

    # ----- Pivot_summary (je Bild + Faktoren + gemittelte Area-Werte) --------
    img_c = sum_h[0]
    ws = wb.create_sheet("Pivot_summary")
    factor_cols = ["animal", "sex", "cond", "panel", "sample"]
    head = [img_c] + factor_cols + sum_h[1:] + (STRUCT_COLS if struct_img else [])
    ws.append(head)
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    numeric_cols = sum_h[1:]
    totals = defaultdict(list)
    for r in summ:
        f = parse_factors(r[img_c])
        if f["is_control"]:
            continue
        si = struct_img.get(r[img_c], {})
        ws.append(
            [r[img_c]] + [f[k] for k in factor_cols] +
            [r.get(h) for h in sum_h[1:]] +
            ([si.get(c) for c in STRUCT_COLS] if struct_img else [])
        )
        for h in numeric_cols:
            if isinstance(r.get(h), float):
                totals[h].append(r[h])
        for c in STRUCT_COLS:
            v = si.get(c)
            if isinstance(v, float):
                totals[c].append(v)
    grow = ["Gesamt/Mittel", "", "", "", "", ""]
    for h in numeric_cols:
        vals = totals.get(h, [])
        grow.append(round(sum(vals), 2) if ("count" in h or h.endswith("vessels") or h == "NK_count") else
                    (round(mean(vals), 3) if vals else None))
    for c in STRUCT_COLS:
        vals = totals.get(c, [])
        grow.append(round(mean(vals), 3) if vals else None)
    ws.append(grow)
    for c in ws[ws.max_row]:
        c.font = HB
    ws.freeze_panes = "A2"

    # ----- Pivot_zones (Zonenprofil je Bild: absolute Zahl + % NK-Zellen je Abstandszone) -----
    zone_pairs = [
        ("NK_count_zone_0_5um",    "pct_zone_0_5um",    "n 0-5µm",  "% 0-5µm"),
        ("NK_count_zone_5_10um",   "pct_zone_5_10um",   "n 5-10µm", "% 5-10µm"),
        ("NK_count_zone_10_20um",  "pct_zone_10_20um",  "n 10-20µm","% 10-20µm"),
        ("NK_count_zone_over20um", "pct_zone_over20um", "n >20µm",  "% >20µm"),
    ]
    zone_pairs = [(nc, pc, nl, pl) for nc, pc, nl, pl in zone_pairs if pc in sum_h]
    if zone_pairs:
        ws = wb.create_sheet("Pivot_zones")
        hdr = [img_c, "animal", "cond", "panel", "n_NK"]
        for _, _, nl, pl in zone_pairs:
            hdr += [nl, pl]
        hdr.append("Summe %")
        ws.append(hdr)
        for c in ws[1]:
            c.font = HB; c.fill = GREY
        ptot = defaultdict(list)
        for r in summ:
            f = parse_factors(r[img_c])
            if f["is_control"]:
                continue
            row = [r[img_c], f["animal"], f["cond"], f["panel"], r.get("NK_count")]
            psum = 0.0
            for nc, pc, _, _ in zone_pairs:
                nv = r.get(nc); pv = r.get(pc)
                row += [int(nv) if isinstance(nv, float) else nv,
                        round(pv, 2) if isinstance(pv, float) else None]
                if isinstance(pv, float):
                    psum += pv; ptot[pc].append(pv)
            row.append(round(psum, 1))
            ws.append(row)
        grow = ["Mittel über Bilder", "", "", "", ""]
        for _, pc, _, _ in zone_pairs:
            vv = ptot.get(pc, [])
            grow += [None, round(mean(vv), 2) if vv else None]
        grow.append(round(sum(g for g in grow[5:] if isinstance(g, (int, float))), 1))
        ws.append(grow)
        for c in ws[ws.max_row]:
            c.font = HB
        ws.freeze_panes = "A2"

    # ----- Pivot_zones_syn (Zonenprofil je Bild: NK-Distanz ZU SYNAPTOPHYSIN) -----
    syn_src = "Dist synaptophysin um"
    if syn_src in det_h:
        img_dd = det_h[0]
        def _syn_zone_idx(d):
            if d <= 5.0:  return 0
            if d <= 10.0: return 1
            if d <= 20.0: return 2
            return 3
        zlabels = [("n 0-5µm", "% 0-5µm"), ("n 5-10µm", "% 5-10µm"),
                   ("n 10-20µm", "% 10-20µm"), ("n >20µm", "% >20µm")]
        zc = defaultdict(lambda: [0, 0, 0, 0])
        for r in det:
            v = r.get(syn_src)
            if isinstance(v, float):
                zc[r[img_dd]][_syn_zone_idx(v)] += 1
        if zc:
            ws = wb.create_sheet("Pivot_zones_syn")
            hdr = [img_dd, "animal", "cond", "panel", "n_NK"]
            for nl, pl in zlabels:
                hdr += [nl, pl]
            hdr.append("Summe %")
            ws.append(hdr)
            for c in ws[1]:
                c.font = HB; c.fill = GREY
            ptot = defaultdict(list)
            for im in sorted(zc):
                f = parse_factors(im)
                if f["is_control"]:
                    continue
                counts = zc[im]; tot = sum(counts)
                row = [im, f["animal"], f["cond"], f["panel"], tot]
                psum = 0.0
                for k, (nl, pl) in enumerate(zlabels):
                    pct = (counts[k] * 100.0 / tot) if tot else None
                    row += [counts[k], round(pct, 2) if pct is not None else None]
                    if pct is not None:
                        psum += pct; ptot[k].append(pct)
                row.append(round(psum, 1))
                ws.append(row)
            grow = ["Mittel über Bilder", "", "", "", ""]
            for k in range(4):
                vv = ptot.get(k, [])
                grow += [None, round(mean(vv), 2) if vv else None]
            grow.append(round(sum(g for g in grow[5:] if isinstance(g, (int, float))), 1))
            ws.append(grow)
            for c in ws[ws.max_row]:
                c.font = HB
            ws.freeze_panes = "A2"

    img_d = det_h[0]
    pick = [
        ("Cell: Area", col_for(det_h, "Cell: Area") and "Cell: Area"),
        ("Cell: Perimeter", "Cell: Perimeter"),
        ("Cell: Circularity", "Cell: Circularity"),
        ("Cell: Max caliper", "Cell: Max caliper"),
        ("Cell: Min caliper", "Cell: Min caliper"),
        ("Cell: NKp46 mean", "Cell: NKp46 mean"),
        ("Cell: NKp46 std dev", "Cell: NKp46 std dev"),
        ("Cell: NKp46 max", "Cell: NKp46 max"),
        ("Dist small vessel um", "Dist small vessel um"),
        ("Dist large vessel um", "Dist large vessel um"),
        ("Dist synaptophysin um", "Dist synaptophysin um"),
        ("Dist nearest vessel um", "Dist nearest vessel um"),
        ("Perivascular (0/1)", "Perivascular (0/1)"),
    ]
    pick = [(lbl, src) for lbl, src in pick if src in det_h]
    agg = defaultdict(lambda: defaultdict(list))
    cnt = defaultdict(int)
    for r in det:
        im = r[img_d]; cnt[im] += 1
        for lbl, src in pick:
            v = r.get(src)
            if isinstance(v, float):
                agg[im][lbl].append(v)
    ws = wb.create_sheet("Pivot_NKcells")
    ws.append([img_d, "animal", "cond", "panel", "n_NK"] + [lbl for lbl, _ in pick])
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    for im in sorted(cnt):
        f = parse_factors(im)
        row = [im, f["animal"], f["cond"], f["panel"], cnt[im]]
        for lbl, _ in pick:
            vals = agg[im][lbl]
            row.append(round(mean(vals), 3) if vals else None)
        ws.append(row)
    ws.freeze_panes = "A2"

    # ----- Pivot_nonNKcells (Distanz-Mittelwerte je Bild, Hintergrund) -----
    if non:
        img_n = non_h[0]
        dist_pick = [(lbl, src) for lbl, src in [
            ("Dist small vessel um", "Dist small vessel um"),
            ("Dist large vessel um", "Dist large vessel um"),
            ("Dist nearest vessel um", "Dist nearest vessel um"),
            ("Dist synaptophysin um", "Dist synaptophysin um"),
            ("Perivascular (0/1)", "Perivascular (0/1)"),
        ] if src in non_h]
        nagg = defaultdict(lambda: defaultdict(list)); ncnt = defaultdict(int)
        for r in non:
            im = r[img_n]; ncnt[im] += 1
            for lbl, src in dist_pick:
                v = r.get(src)
                if isinstance(v, float):
                    nagg[im][lbl].append(v)
        ws = wb.create_sheet("Pivot_nonNKcells")
        ws.append([img_n, "animal", "cond", "panel", "n_nonNK"] + [lbl for lbl, _ in dist_pick])
        for c in ws[1]:
            c.font = HB; c.fill = GREY
        for im in sorted(ncnt):
            f = parse_factors(im)
            row = [im, f["animal"], f["cond"], f["panel"], ncnt[im]]
            for lbl, _ in dist_pick:
                vals = nagg[im][lbl]
                row.append(round(mean(vals), 3) if vals else None)
            ws.append(row)
        ws.freeze_panes = "A2"

    # ----- Pivot_vessels (Summen je Bild) -----
    vcols = [c for c in ["Num Detections", "Num NK-cell", "NK within 20um count", "NK within 5um count"] if c in ves_h]
    vagg = defaultdict(lambda: defaultdict(float)); vcnt = defaultdict(int)
    for r in ves:
        im = r[img_v]; vcnt[im] += 1
        for c in vcols:
            v = r.get(c)
            if isinstance(v, float):
                vagg[im][c] += v
    ws = wb.create_sheet("Pivot_vessels")
    ws.append([img_v, "animal", "cond", "panel", "n_annot"] + [f"Summe {c}" for c in vcols])
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    for im in sorted(vcnt):
        f = parse_factors(im)
        ws.append([im, f["animal"], f["cond"], f["panel"], vcnt[im]] + [round(vagg[im][c], 1) for c in vcols])
    ws.freeze_panes = "A2"

    # ----- Pivot_struct_area (je Bild: Dichte & Fläche der Strukturen) -----
    if struct_raw:
        ws = wb.create_sheet("Pivot_struct_area")
        ws.append([img_v, "animal", "cond", "panel",
                   "tissue_area_mm2",
                   "n_synaptophysin", "syn_density_per_mm2", "syn_area_mean_um2", "syn_area_fraction_pct",
                   "n_small_vessel", "small_vessel_area_mean_um2",
                   "n_large_vessel", "large_vessel_area_mean_um2"])
        for c in ws[1]:
            c.font = HB; c.fill = GREY
        for im in sorted(struct_raw):
            f = parse_factors(im)
            if f["is_control"]:
                continue
            s = struct_raw[im]; tmm2 = s["tissue_um2"] / 1e6 if s["tissue_um2"] else None
            syn, sm, lg = s["syn"], s["small"], s["large"]
            ws.append([
                im, f["animal"], f["cond"], f["panel"],
                round(tmm2, 4) if tmm2 else None,
                len(syn),
                round(len(syn) / tmm2, 2) if (syn and tmm2) else None,
                round(mean(syn), 2) if syn else None,
                round(sum(syn) * 100.0 / s["tissue_um2"], 4) if (syn and s["tissue_um2"]) else None,
                len(sm),
                round(mean(sm), 2) if sm else None,
                len(lg),
                round(mean(lg), 2) if lg else None,
            ])
        ws.freeze_panes = "A2"

    # ----- PerAnimal_condition (inkl. gemittelter Area-Werte) ----------------
    metrics = [c for c in ["NK_density_per_mm2",
                           "small_vessel_density_per_mm2", "large_vessel_density_per_mm2", "vessel_density_per_mm2",
                           "perivascular_pct", "vessel_area_fraction_pct",
                           "mean_dist_small_um", "mean_dist_large_um", "mean_dist_syn_um",
                           "NK_count_zone_0_5um", "pct_zone_0_5um",
                           "NK_count_zone_5_10um", "pct_zone_5_10um",
                           "NK_count_zone_10_20um", "pct_zone_10_20um",
                           "NK_count_zone_over20um", "pct_zone_over20um"] if c in sum_h]
    pa = defaultdict(list)
    for r in summ:
        f = parse_factors(r[img_c])
        if f["is_control"]:
            continue
        for m in metrics:
            if isinstance(r.get(m), float):
                pa[(f["animal"], f["cond"], m)].append(r[m])
    # Struktur-Area-Mittel je Tier x Bedingung
    pa_struct = defaultdict(list)
    for _im, _si in struct_img.items():
        _f = parse_factors(_im)
        if _f["is_control"]:
            continue
        for c in STRUCT_COLS:
            v = _si.get(c)
            if v is not None:
                pa_struct[(_f["animal"], _f["cond"], c)].append(v)
    ws = wb.create_sheet("PerAnimal_condition")
    ws.append(["animal", "cond"] + metrics + (STRUCT_COLS if struct_img else []))
    for c in ws[1]:
        c.font = HB; c.fill = GREY
    animals = sorted({parse_factors(r[img_c])["animal"] for r in summ if not parse_factors(r[img_c])["is_control"]})
    for a in animals:
        for cond in ["OP", "N"]:
            row = [a, cond]
            for m in metrics:
                vals = pa.get((a, cond, m), [])
                row.append(round(mean(vals), 2) if vals else None)
            for c in STRUCT_COLS:
                vals = pa_struct.get((a, cond, c), [])
                row.append(round(mean(vals), 2) if vals else None)
            ws.append(row)
    ws.freeze_panes = "A2"

    # ----- Grunddaten (Rohkopien) -----
    write_raw(wb, "Grunddaten_summary", sum_h, summ)
    write_raw(wb, "Grunddaten_vessels", ves_h, ves)
    write_raw(wb, "Grunddaten_NKcells", det_h, det)
    if non:
        write_raw(wb, "Grunddaten_nonNK", non_h, non)

    wb.save(out)
    print(f"WROTE {out}  ({len(wb.sheetnames)} sheets)")

if __name__ == "__main__":
    measdir = sys.argv[1] if len(sys.argv) > 1 else "."
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(measdir, "QuPath_Auswertung.xlsx")
    build(measdir, out)

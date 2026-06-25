#!/usr/bin/env python3
"""
QuPath-Exporte -> Prism-fertige Tabellen (Spaltenbloecke zum direkten Einfuegen).
Liest ALL_detections_NKcells / ALL_vessels_annotations / SUMMARY_per_image (TSV/CSV)
und schreibt Prism_ready_analysis.xlsx.

Aufruf:  python make_prism_tables.py  <measurements_dir>  [output.xlsx]
"""
import sys, os, csv, re, glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from collections import defaultdict
from statistics import mean

HB = Font(bold=True); H1 = Font(bold=True, size=13); GREY = PatternFill("solid", fgColor="DDDDDD")

def find(measdir, stem):
    for ext in (".tsv", ".csv"):
        p = os.path.join(measdir, stem + ext)
        if os.path.exists(p): return p
    h = glob.glob(os.path.join(measdir, stem + "*"))
    return h[0] if h else None

def load(path):
    if path is None: return [], []
    with open(path, newline="", encoding="utf-8-sig") as f:
        s = f.read(4096); f.seek(0)
        d = "\t" if s.count("\t") >= max(s.count(";"), s.count(",")) else (";" if s.count(";") >= s.count(",") else ",")
        rd = csv.reader(f, delimiter=d); hdr = next(rd); rows = []
        for r in rd:
            o = {}
            for k, v in zip(hdr, r):
                if v is None or v == "": o[k] = None
                else:
                    try: o[k] = float(v.replace(",", ".")) if re.match(r"^-?\d", v) else v
                    except ValueError: o[k] = v
            rows.append(o)
    return hdr, rows

def parse(img):
    s = str(img)
    return dict(
        animal=(re.search(r"_(\d{4,6})_", s) or [None, None])[1],
        cond="OP" if "gcOP" in s else ("N" if "gcN" in s else "?"),
        panel="IB4" if "IB4" in s else ("Syn" if "Syn" in s else "?"),
        ctrl="control" in s.lower())

def main(measdir, out):
    det_h, det = load(find(measdir, "ALL_detections_NKcells"))
    ves_h, ves = load(find(measdir, "ALL_vessels_annotations"))
    sum_h, summ = load(find(measdir, "SUMMARY_per_image"))
    recs = []
    for r in summ:
        f = parse(r["Image"])
        if f["ctrl"]: continue
        f.update(r); f["image"] = r["Image"]; recs.append(f)

    # --- Struktur-Flaechen & -Dichte aus den Annotationen (small/large vessel, Synaptophysin) ---
    # Spalten robust gegen µm-Encoding aufloesen.
    area_col = next((h for h in ves_h if "area" in h.lower()), None)
    cls_col  = next((h for h in ves_h if h.strip().lower() == "classification"), None)
    img_col  = ves_h[0] if ves_h else "Image"
    struct = defaultdict(lambda: {"tissue_um2": 0.0, "syn": [], "small": [], "large": []})
    if area_col and cls_col:
        for r in ves:
            cls = (r.get(cls_col) or "").strip(); a = r.get(area_col); im = r.get(img_col)
            if not isinstance(a, float):
                continue
            if   cls == "Tissue":        struct[im]["tissue_um2"] += a
            elif cls == "Synaptophysin": struct[im]["syn"].append(a)
            elif cls == "small vessel":  struct[im]["small"].append(a)
            elif cls == "large vessel":  struct[im]["large"].append(a)
    # je Bild abgeleitete Kennzahlen an die Records anhaengen
    for d in recs:
        s = struct.get(d["image"])
        if not s:
            continue
        tmm2 = s["tissue_um2"] / 1e6 if s["tissue_um2"] else None
        if s["syn"]:
            d["syn_struct_area_mean"]  = mean(s["syn"])
            if tmm2:
                d["syn_struct_density"] = len(s["syn"]) / tmm2
            if s["tissue_um2"]:
                d["syn_area_fraction_pct"] = sum(s["syn"]) * 100.0 / s["tissue_um2"]
        if s["small"]:
            d["small_vessel_area_mean"] = mean(s["small"])
        if s["large"]:
            d["large_vessel_area_mean"] = mean(s["large"])

    wb = Workbook(); ws0 = wb.active; ws0.title = "README"
    for i, (t, fnt) in enumerate([
        ("Prism-fertige Tabellen – welches Sheet in welche Prism-Analyse", H1),
        ("Jeder Block: Kopfzeile = Gruppen, darunter je Bild ein Wert. In Prism einfuegen.", None),
        ("BIOL. REPLIKAT = TIER (hier n=2) -> p-Werte deskriptiv/Pilot! AnimalLevel_means zeigt Tiermittel.", HB),
        ("", None),
        ("A1_NKdensity_*  : NK-Dichte OP vs N (je Panel)        -> Mann-Whitney", None),
        ("A2_Perivascular : % NK <20µm am Gefaess, OP vs N       -> Mann-Whitney", None),
        ("A3_Dist_small_vs_large : gepaart je IB4-Bild           -> Wilcoxon paired", None),
        ("A4_Dist_to_Syn  : NK->Synaptophysin OP vs N            -> Mann-Whitney", None),
        ("A5_Caliber_NKwithin20 : small vs large, gepaart        -> Wilcoxon paired", None),
        ("A6_Activation_by_location : NKp46/Flaeche peri vs nicht -> Wilcoxon paired", None),
        ("A7_Perivasc_NK_vs_background : NK vs Nicht-NK, gepaart  -> Wilcoxon paired (ANREICHERUNG)", HB),
        ("A8_DistSyn_NK_vs_background  : NK vs Nicht-NK, gepaart  -> Wilcoxon paired (ANREICHERUNG)", HB),
        ("", None),
        ("A9_Zone_0_5um   : % NK ≤5µm vom Gefäß, OP vs N            -> Mann-Whitney", None),
        ("A10_Zone_5_10um : % NK 5-10µm vom Gefäß, OP vs N          -> Mann-Whitney", None),
        ("A11_Zone_10_20um: % NK 10-20µm vom Gefäß, OP vs N         -> Mann-Whitney", None),
        ("A12_Zone_over20um: % NK >20µm vom Gefäß, OP vs N          -> Mann-Whitney", None),
        ("", None),
        ("A13_SmallVesselDensity : small vessels/mm², OP vs N       -> Mann-Whitney", None),
        ("A14_LargeVesselDensity : large vessels/mm², OP vs N       -> Mann-Whitney", None),
        ("", None),
        ("A15_ZoneProfile_perImage : % NK je Zone, Zonen als Zeilen, OP/N -> Grouped (Verteilung)", HB),
        ("A16_ZoneProfile_perAnimal: % NK je Zone, Tier-Mittel (n=2)       -> Grouped (Verteilung)", HB),
        ("", None),
        ("--- SYNAPTOPHYSIN-DISTANZZONEN (NK -> Synaptophysin) ---", HB),
        ("A17_SynZone_0_5um   : % NK ≤5µm zu Synaptophysin, OP vs N  -> Mann-Whitney", None),
        ("A18_SynZone_5_10um  : % NK 5-10µm zu Synaptophysin, OP vs N -> Mann-Whitney", None),
        ("A19_SynZone_10_20um : % NK 10-20µm zu Synaptophysin, OP vs N-> Mann-Whitney", None),
        ("A20_SynZone_over20um: % NK >20µm zu Synaptophysin, OP vs N  -> Mann-Whitney", None),
        ("A21_SynZoneProfile_perImage : % NK je Syn-Zone, je Bild     -> Grouped (Verteilung)", HB),
        ("A22_SynZoneProfile_perAnimal: % NK je Syn-Zone, Tier-Mittel -> Grouped (Verteilung)", HB),
        ("", None),
        ("--- STRUKTUR-DICHTE & -FLAECHE (Synaptophysin / Gefaesse) ---", HB),
        ("A23_SynStructDensity : Synaptophysin-Strukturen/mm², OP vs N -> Mann-Whitney", None),
        ("A24_SynStructArea    : Ø Synaptophysin-Flaeche µm², OP vs N  -> Mann-Whitney", None),
        ("A24b_SynAreaFraction : Synaptophysin-Flaechenanteil %, OP vs N-> Mann-Whitney", None),
        ("A25_SmallVesselArea  : Ø small-vessel-Flaeche µm², OP vs N   -> Mann-Whitney", None),
        ("A26_LargeVesselArea  : Ø large-vessel-Flaeche µm², OP vs N   -> Mann-Whitney", None),
        ("A27_VesselArea_small_vs_large : Ø Flaeche gepaart small/large -> Wilcoxon paired", None),
    ], 1):
        c = ws0.cell(i, 1, t)
        if fnt: c.font = fnt
    ws0.column_dimensions['A'].width = 95

    def block(title, groups, data, note):
        ws = wb.create_sheet(title); ws.cell(1, 1, note).font = HB
        for j, g in enumerate(groups, 1): ws.cell(3, j, g).font = HB; ws.cell(3, j).fill = GREY
        n = max((len(data[g]) for g in groups), default=0)
        for i in range(n):
            for j, g in enumerate(groups, 1):
                if i < len(data[g]) and data[g][i] is not None:
                    ws.cell(4 + i, j, round(float(data[g][i]), 3))
        for j in range(len(groups)): ws.column_dimensions[chr(65 + j)].width = 16

    def paired(title, cols, getrow, note, src):
        ws = wb.create_sheet(title); ws.cell(1, 1, note).font = HB
        for j, h in enumerate(cols, 1): ws.cell(3, j, h).font = HB; ws.cell(3, j).fill = GREY
        ri = 4
        for d in src:
            row = getrow(d)
            if row is None: continue
            for j, v in enumerate(row, 1):
                ws.cell(ri, j, round(v, 3) if isinstance(v, float) else v)
            ri += 1
        ws.column_dimensions['A'].width = 42

    # A1 density per panel
    for panel in ["IB4", "Syn"]:
        g = {"OP": [], "N": []}
        for d in recs:
            if d["panel"] == panel and isinstance(d.get("NK_density_per_mm2"), float):
                g[d["cond"]].append(d["NK_density_per_mm2"])
        block(f"A1_NKdensity_{panel}", ["OP", "N"], g, f"NK-Dichte (Zellen/mm²), {panel}. Mann-Whitney.")

    # A2 perivascular OP vs N
    g = {"OP": [], "N": []}
    for d in recs:
        if d["panel"] == "IB4" and isinstance(d.get("perivascular_pct"), float):
            g[d["cond"]].append(d["perivascular_pct"])
    block("A2_Perivascular_pct", ["OP", "N"], g, "% NK <20µm am Gefaess, IB4. Mann-Whitney.")

    # A3 small vs large paired
    paired("A3_Dist_small_vs_large", ["Bild", "cond", "small_um", "large_um"],
           lambda d: ([d["image"][:40], d["cond"], d["mean_dist_small_um"], d["mean_dist_large_um"]]
                      if d["panel"] == "IB4" and isinstance(d.get("mean_dist_small_um"), float)
                      and isinstance(d.get("mean_dist_large_um"), float) else None),
           "Mittl. Distanz NK->Gefaess, gepaart small vs large je IB4-Bild. Wilcoxon paired.", recs)

    # A4 dist to syn OP vs N
    g = {"OP": [], "N": []}
    for d in recs:
        if d["panel"] == "Syn" and isinstance(d.get("mean_dist_syn_um"), float):
            g[d["cond"]].append(d["mean_dist_syn_um"])
    block("A4_Dist_to_Syn", ["OP", "N"], g, "Mittl. Distanz NK->Synaptophysin [µm], Syn. Mann-Whitney.")

    # A5 caliber from vessels
    cal = defaultdict(lambda: defaultdict(list))
    for r in ves:
        cls = r.get("Classification")
        if cls not in ("small vessel", "large vessel"): continue
        w = r.get("NK within 20um count")
        if isinstance(w, float): cal[r["Image"]][cls].append(w)
    ws = wb.create_sheet("A5_Caliber_NKwithin20"); ws.cell(1, 1, "Mittl. NK im 20µm-Umkreis je Gefaess; gepaart small vs large. Wilcoxon paired.").font = HB
    for j, h in enumerate(["Bild", "small_meanNK", "large_meanNK"], 1): ws.cell(3, j, h).font = HB; ws.cell(3, j).fill = GREY
    ri = 4
    for im, cc in cal.items():
        if "small vessel" in cc and "large vessel" in cc:
            ws.cell(ri, 1, im[:40]); ws.cell(ri, 2, round(mean(cc["small vessel"]), 3)); ws.cell(ri, 3, round(mean(cc["large vessel"]), 3)); ri += 1
    ws.column_dimensions['A'].width = 42

    # A6 activation by location (per IB4 image)
    peri_nk = defaultdict(list); non_nk = defaultdict(list); peri_a = defaultdict(list); non_a = defaultdict(list)
    for r in det:
        p = r.get("Perivascular (0/1)")
        if p is None: continue
        im = r["Image"]; nk = r.get("Cell: NKp46 mean"); ar = r.get("Cell: Area")
        (peri_nk if p == 1 else non_nk)[im].append(nk) if isinstance(nk, float) else None
        (peri_a if p == 1 else non_a)[im].append(ar) if isinstance(ar, float) else None
    ws = wb.create_sheet("A6_Activation_by_location"); ws.cell(1, 1, "Pro IB4-Bild: NKp46 & Zellflaeche, perivaskulaer vs nicht. Wilcoxon paired.").font = HB
    for j, h in enumerate(["Bild", "cond", "NKp46_peri", "NKp46_nonPeri", "Area_peri", "Area_nonPeri"], 1):
        ws.cell(3, j, h).font = HB; ws.cell(3, j).fill = GREY
    ri = 4
    for d in recs:
        if d["panel"] != "IB4": continue
        im = d["image"]
        if im in peri_nk and im in non_nk and peri_nk[im] and non_nk[im]:
            ws.cell(ri, 1, im[:40]); ws.cell(ri, 2, d["cond"])
            ws.cell(ri, 3, round(mean(peri_nk[im]), 1)); ws.cell(ri, 4, round(mean(non_nk[im]), 1))
            ws.cell(ri, 5, round(mean(peri_a[im]), 2)); ws.cell(ri, 6, round(mean(non_a[im]), 2)); ri += 1
    ws.column_dimensions['A'].width = 42

    # A7 enrichment perivascular NK vs background
    paired("A7_Perivasc_NK_vs_background", ["Bild", "cond", "NK_perivasc", "nonNK_perivasc"],
           lambda d: ([d["image"][:40], d["cond"], d["perivascular_pct"], d["nonNK_perivascular_pct"]]
                      if d["panel"] == "IB4" and isinstance(d.get("perivascular_pct"), float)
                      and isinstance(d.get("nonNK_perivascular_pct"), float) else None),
           "ANREICHERUNG: % perivaskulaer NK vs Nicht-NK, gepaart je IB4-Bild. Wilcoxon paired.", recs)

    # A8 enrichment dist to syn NK vs background
    paired("A8_DistSyn_NK_vs_background", ["Bild", "cond", "NK_dist_syn", "nonNK_dist_syn"],
           lambda d: ([d["image"][:40], d["cond"], d["mean_dist_syn_um"], d["nonNK_mean_dist_syn_um"]]
                      if d["panel"] == "Syn" and isinstance(d.get("mean_dist_syn_um"), float)
                      and isinstance(d.get("nonNK_mean_dist_syn_um"), float) else None),
           "ANREICHERUNG: NK vs Nicht-NK Distanz->Synaptophysin, gepaart je Syn-Bild. Wilcoxon paired.", recs)

    # A9-A12 distance zones (IB4, OP vs N)
    for sheet, col, label in [
        ("A9_Zone_0_5um",    "pct_zone_0_5um",    "≤5µm"),
        ("A10_Zone_5_10um",  "pct_zone_5_10um",   "5–10µm"),
        ("A11_Zone_10_20um", "pct_zone_10_20um",  "10–20µm"),
        ("A12_Zone_over20um","pct_zone_over20um",  ">20µm"),
    ]:
        g = {"OP": [], "N": []}
        for d in recs:
            if d["panel"] == "IB4" and isinstance(d.get(col), float):
                g[d["cond"]].append(d[col])
        block(sheet, ["OP", "N"], g, f"% NK im Abstand {label} vom nächsten Gefäß, IB4. Mann-Whitney.")

    # A13/A14 vessel density (IB4, OP vs N)
    for sheet, col, label in [
        ("A13_SmallVesselDensity", "small_vessel_density_per_mm2", "Small vessels"),
        ("A14_LargeVesselDensity", "large_vessel_density_per_mm2", "Large vessels"),
    ]:
        g = {"OP": [], "N": []}
        for d in recs:
            if d["panel"] == "IB4" and isinstance(d.get(col), float):
                g[d["cond"]].append(d[col])
        block(sheet, ["OP", "N"], g, f"{label} pro mm² Gewebe, IB4. Mann-Whitney.")

    # A15/A16 Zonenprofil: n (absolut) und % NK je Abstandszone -> Prism Grouped
    zone_pairs_prism = [
        ("0-5µm",   "NK_count_zone_0_5um",    "pct_zone_0_5um"),
        ("5-10µm",  "NK_count_zone_5_10um",   "pct_zone_5_10um"),
        ("10-20µm", "NK_count_zone_10_20um",  "pct_zone_10_20um"),
        (">20µm",   "NK_count_zone_over20um", "pct_zone_over20um"),
    ]
    zone_pairs_prism = [(lbl, nc, pc) for lbl, nc, pc in zone_pairs_prism
                        if any(pc in r for r in recs)]
    # Für Rückwärtskompatibilität mit alten Exports (nur pct, kein count)
    zones = [(lbl, pc) for lbl, nc, pc in zone_pairs_prism]

    def zone_grouped(title, op_keys, n_keys, getval_n, getval_p, note, zpairs=None):
        zp = zpairs if zpairs is not None else zone_pairs_prism
        ws = wb.create_sheet(title); ws.cell(1, 1, note).font = HB
        kmax = max(len(op_keys), len(n_keys), 1)
        # Kopfzeile: n-OP ... n-N ... %-OP ... %-N
        ws.cell(3, 1, "Zone").font = HB; ws.cell(3, 1).fill = GREY
        base_n_op = 2; base_n_n = 2 + kmax
        base_p_op = 2 + 2*kmax; base_p_n = 2 + 3*kmax
        for j in range(kmax):
            for base, lbl in [(base_n_op,"n OP"),(base_n_n,"n N"),(base_p_op,"% OP"),(base_p_n,"% N")]:
                c = ws.cell(3, base + j, lbl); c.font = HB; c.fill = GREY
        for i, (zlbl, nc, pc) in enumerate(zp):
            r = 4 + i
            ws.cell(r, 1, zlbl).font = HB
            for j, key in enumerate(op_keys):
                nv = getval_n(key, "OP", nc); pv = getval_p(key, "OP", pc)
                if nv is not None: ws.cell(r, base_n_op + j, round(nv, 1))
                if pv is not None: ws.cell(r, base_p_op + j, round(pv, 2))
            for j, key in enumerate(n_keys):
                nv = getval_n(key, "N", nc); pv = getval_p(key, "N", pc)
                if nv is not None: ws.cell(r, base_n_n + j, round(nv, 1))
                if pv is not None: ws.cell(r, base_p_n + j, round(pv, 2))
        ws.column_dimensions['A'].width = 22

    if zone_pairs_prism:
        # A15: je IB4-Bild
        op_imgs = [d["image"] for d in recs if d["panel"] == "IB4" and d["cond"] == "OP"]
        n_imgs  = [d["image"] for d in recs if d["panel"] == "IB4" and d["cond"] == "N"]
        by_img  = {d["image"]: d for d in recs if d["panel"] == "IB4"}
        def gv_img_n(img, cond, col):
            d = by_img.get(img); v = d.get(col) if d else None
            return float(v) if v is not None and not isinstance(v, str) else None
        zone_grouped("A15_ZoneProfile_perImage", op_imgs, n_imgs, gv_img_n, gv_img_n,
                     "NK je Zone: n (absolut) und % je IB4-Bild. Replikat = Bild. Prism: Grouped.")

        # A16: Tier-Mittel
        animals = sorted({d["animal"] for d in recs if d["panel"] == "IB4"})
        amean = defaultdict(list)
        for d in recs:
            if d["panel"] != "IB4": continue
            for _, nc, pc in zone_pairs_prism:
                for col in (nc, pc):
                    v = d.get(col)
                    if v is not None and not isinstance(v, str):
                        amean[(d["animal"], d["cond"], col)].append(float(v))
        def gv_animal(animal, cond, col):
            vv = amean.get((animal, cond, col)); return mean(vv) if vv else None
        zone_grouped("A16_ZoneProfile_perAnimal", animals, animals, gv_animal, gv_animal,
                     "NK je Zone: n und % Tier-Mittel (n=2, biolog. Replikat). Prism: Grouped.")

    # =====================================================================
    # A17-A22  SYNAPTOPHYSIN-DISTANZZONEN (NK-Zelle -> naechste Synaptophysin-
    # Annotation). Pro NK-Zelle aus der Rohdistanz 'Dist synaptophysin um'
    # berechnet, gleiche Zonengrenzen wie bei den Gefaessen (<=5 / <=10 / <=20 / >20).
    # Panel = Syn (nur Synaptophysin-gefaerbte Schnitte).
    # =====================================================================
    syn_src = "Dist synaptophysin um"
    if syn_src in det_h:
        def _syn_idx(d):
            if d <= 5.0:  return 0
            if d <= 10.0: return 1
            if d <= 20.0: return 2
            return 3
        syn_cnt = defaultdict(lambda: [0, 0, 0, 0])
        for r in det:
            v = r.get(syn_src)
            if isinstance(v, float):
                syn_cnt[r["Image"]][_syn_idx(v)] += 1
        # n und % je Zone an die (Syn-)Bild-Records anhaengen
        for d in recs:
            c = syn_cnt.get(d["image"])
            if not c:
                continue
            tot = sum(c)
            for k in range(4):
                d[f"syn_count_zone_{k}"] = c[k]
                d[f"syn_pct_zone_{k}"]   = (c[k] * 100.0 / tot) if tot else None

        # A17-A20: % NK je Syn-Zone, OP vs N (Mann-Whitney)
        for sheet, k, label in [
            ("A17_SynZone_0_5um",    0, "≤5µm"),
            ("A18_SynZone_5_10um",   1, "5–10µm"),
            ("A19_SynZone_10_20um",  2, "10–20µm"),
            ("A20_SynZone_over20um", 3, ">20µm"),
        ]:
            col = f"syn_pct_zone_{k}"
            g = {"OP": [], "N": []}
            for d in recs:
                if d["panel"] == "Syn" and isinstance(d.get(col), float):
                    g[d["cond"]].append(d[col])
            block(sheet, ["OP", "N"], g,
                  f"% NK im Abstand {label} zur naechsten Synaptophysin-Annotation, Syn. Mann-Whitney.")

        # A21/A22: Zonenprofil (n + %) je Syn-Bild bzw. Tier-Mittel -> Prism Grouped
        syn_zone_pairs = [
            ("0-5µm",   "syn_count_zone_0", "syn_pct_zone_0"),
            ("5-10µm",  "syn_count_zone_1", "syn_pct_zone_1"),
            ("10-20µm", "syn_count_zone_2", "syn_pct_zone_2"),
            (">20µm",   "syn_count_zone_3", "syn_pct_zone_3"),
        ]
        syn_recs = [d for d in recs if d["panel"] == "Syn" and "syn_count_zone_0" in d]
        if syn_recs:
            # A21: je Syn-Bild
            op_imgs_s = [d["image"] for d in syn_recs if d["cond"] == "OP"]
            n_imgs_s  = [d["image"] for d in syn_recs if d["cond"] == "N"]
            by_img_s  = {d["image"]: d for d in syn_recs}
            def gv_simg(img, cond, col):
                d = by_img_s.get(img); v = d.get(col) if d else None
                return float(v) if v is not None and not isinstance(v, str) else None
            zone_grouped("A21_SynZoneProfile_perImage", op_imgs_s, n_imgs_s, gv_simg, gv_simg,
                         "NK je Synaptophysin-Zone: n (absolut) und % je Syn-Bild. Replikat = Bild. Prism: Grouped.",
                         zpairs=syn_zone_pairs)

            # A22: Tier-Mittel
            animals_s = sorted({d["animal"] for d in syn_recs})
            amean_s = defaultdict(list)
            for d in syn_recs:
                for _, nc, pc in syn_zone_pairs:
                    for col in (nc, pc):
                        v = d.get(col)
                        if v is not None and not isinstance(v, str):
                            amean_s[(d["animal"], d["cond"], col)].append(float(v))
            def gv_sanimal(animal, cond, col):
                vv = amean_s.get((animal, cond, col)); return mean(vv) if vv else None
            zone_grouped("A22_SynZoneProfile_perAnimal", animals_s, animals_s, gv_sanimal, gv_sanimal,
                         "NK je Synaptophysin-Zone: n und % Tier-Mittel (biolog. Replikat). Prism: Grouped.",
                         zpairs=syn_zone_pairs)

    # =====================================================================
    # A23-A27  STRUKTUR-DICHTE & -FLAECHE (Synaptophysin / small / large vessel)
    # je Bild aus den Annotationen berechnet, OP vs N. Replikat = Bild.
    # =====================================================================
    # A23: Synaptophysin-Dichte (Strukturen pro mm² Gewebe), Syn-Panel
    g = {"OP": [], "N": []}
    for d in recs:
        if d["panel"] == "Syn" and isinstance(d.get("syn_struct_density"), float):
            g[d["cond"]].append(d["syn_struct_density"])
    block("A23_SynStructDensity", ["OP", "N"], g,
          "Synaptophysin-Strukturen pro mm² Gewebe, Syn. Mann-Whitney.")

    # A24: mittlere Synaptophysin-Strukturflaeche [µm²], Syn-Panel
    g = {"OP": [], "N": []}
    for d in recs:
        if d["panel"] == "Syn" and isinstance(d.get("syn_struct_area_mean"), float):
            g[d["cond"]].append(d["syn_struct_area_mean"])
    block("A24_SynStructArea", ["OP", "N"], g,
          "Mittlere Synaptophysin-Strukturflaeche [µm²], Syn. Mann-Whitney.")

    # A24b: Synaptophysin-Flaechenanteil am Gewebe [%], Syn-Panel
    g = {"OP": [], "N": []}
    for d in recs:
        if d["panel"] == "Syn" and isinstance(d.get("syn_area_fraction_pct"), float):
            g[d["cond"]].append(d["syn_area_fraction_pct"])
    block("A24b_SynAreaFraction", ["OP", "N"], g,
          "Synaptophysin-Flaechenanteil am Gewebe [%], Syn. Mann-Whitney.")

    # A25/A26: mittlere Gefaessflaeche [µm²], small bzw. large, IB4-Panel, OP vs N
    for sheet, key, label in [
        ("A25_SmallVesselArea", "small_vessel_area_mean", "Small vessels"),
        ("A26_LargeVesselArea", "large_vessel_area_mean", "Large vessels"),
    ]:
        g = {"OP": [], "N": []}
        for d in recs:
            if d["panel"] == "IB4" and isinstance(d.get(key), float):
                g[d["cond"]].append(d[key])
        block(sheet, ["OP", "N"], g, f"Mittlere {label}-Flaeche [µm²], IB4. Mann-Whitney.")

    # A27: small vs large Gefaessflaeche, gepaart je IB4-Bild
    paired("A27_VesselArea_small_vs_large", ["Bild", "cond", "small_area_um2", "large_area_um2"],
           lambda d: ([d["image"][:40], d["cond"], d["small_vessel_area_mean"], d["large_vessel_area_mean"]]
                      if d["panel"] == "IB4" and isinstance(d.get("small_vessel_area_mean"), float)
                      and isinstance(d.get("large_vessel_area_mean"), float) else None),
           "Mittlere Strukturflaeche [µm²], gepaart small vs large je IB4-Bild. Wilcoxon paired.", recs)

    # AnimalLevel means
    metrics = ["NK_density_per_mm2",
               "small_vessel_density_per_mm2", "large_vessel_density_per_mm2",
               "perivascular_pct",
               "NK_count_zone_0_5um",   "pct_zone_0_5um",
               "NK_count_zone_5_10um",  "pct_zone_5_10um",
               "NK_count_zone_10_20um", "pct_zone_10_20um",
               "NK_count_zone_over20um","pct_zone_over20um",
               "mean_dist_small_um", "mean_dist_large_um",
               "mean_dist_syn_um", "nonNK_perivascular_pct", "nonNK_mean_dist_syn_um"]
    metrics = [m for m in metrics if m in sum_h]
    # Struktur-Dichte/-Flaeche (aus Annotationen berechnet, nicht in SUMMARY) ergaenzen
    struct_metrics = ["syn_struct_density", "syn_struct_area_mean", "syn_area_fraction_pct",
                      "small_vessel_area_mean", "large_vessel_area_mean"]
    struct_metrics = [m for m in struct_metrics if any(isinstance(d.get(m), float) for d in recs)]
    metrics = metrics + struct_metrics
    agg = defaultdict(list)
    for d in recs:
        for m in metrics:
            if isinstance(d.get(m), float): agg[(d["animal"], d["cond"], m)].append(d[m])
    ws = wb.create_sheet("AnimalLevel_means"); ws.cell(1, 1, "Tier-Mittel (richtiges biol. Replikat, n=2).").font = HB
    for j, h in enumerate(["animal", "cond"] + metrics, 1): ws.cell(3, j, h).font = HB; ws.cell(3, j).fill = GREY
    ri = 4
    for a in sorted({d["animal"] for d in recs}):
        for cond in ["OP", "N"]:
            ws.cell(ri, 1, a); ws.cell(ri, 2, cond)
            for j, m in enumerate(metrics, 3):
                v = agg.get((a, cond, m))
                if v: ws.cell(ri, j, round(mean(v), 2))
            ri += 1
    for j in range(len(metrics) + 2): ws.column_dimensions[chr(65 + j)].width = 18

    wb.save(out); print("WROTE", out, "(", len(wb.sheetnames), "sheets )")

if __name__ == "__main__":
    md = sys.argv[1] if len(sys.argv) > 1 else "."
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(md, "Prism_ready_analysis.xlsx")
    main(md, out)
